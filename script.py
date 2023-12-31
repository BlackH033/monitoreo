import rasterio as rs
import matplotlib.pyplot as plt
import numpy as np
import os
from windows import *
from shapely.geometry import shape
import geopandas as gpd
from datetime import datetime
from rasterio.features import geometry_mask
import openpyxl
from rasterio.warp import calculate_default_transform, reproject
from pyproj import Transformer
import numpy as np
#--------------------------------------------------
class procesamiento():
    carpeta_raiz=os.path.dirname(__file__)          #guarda la ruta donde se encuentra este archivo .py
    carpeta_img=os.path.join(carpeta_raiz,"img")    #crea la ruta relativa a la carpeta /img - la cual se guarda en la misma ruta del archivo .py
    ruta=""
    def __init__(self,ruta,tipo):
        super().__init__()
        self.ruta=ruta
        self.directorio=os.listdir(ruta)
        print(self.ruta)
        #funciones

        #ndvi=(nir-red)/(nir+red)
        def ndvi(red,nir):
            #ndvi=np.where((nir+red)==0,0,(nir-red)/(nir+red))
            return (nir-red)/(nir+red)

        def verdadero_color(r,g,b):
            r=r/r.max()
            g=g/g.max()
            b=b/b.max()
            rgb=np.dstack((r,g,b))
            return rgb

        def porcentaje(x,años):
            if len(x)==1:return "Dato inicial"
            else:
                porc=(x[-1]-x[-2])/x[-2]*100
                if porc>0:return f"Incremento del {round(porc,3)}% con\nrespecto a{años[-2]}"
                else:return f"Decremento del {round(abs(porc),3)}% con\nrespecto a {años[-2]}"

        def info(ax,x,m):
            ax.text(0.5, 0.95, f"Pixeles totales: {x.size}", ha='center', va='center', fontsize=12, color='black')
            ax.text(0.5, 0.9, f"Dimensión: {x.shape}", ha='center', va='center', fontsize=12, color='black')
            ax.text(0.5, 0.8, f"Valor NDVI minimo: {round(np.nanmin(x),3)}", ha='center', va='center', fontsize=12, color='black')
            ax.text(0.5, 0.75, f"Valor NDVI maximo: {round(np.nanmax(x),3)}", ha='center', va='center', fontsize=12, color='black')
            ax.text(0.5, 0.7, f"Valor NDVI promedio: {round(np.nanmean(x),3)}", ha='center', va='center', fontsize=12, color='black')
            ax.text(0.5, 0.65, f"_______________________________________", ha='center', va='center', fontsize=12, color='black')
            ax.text(0.5, 0.5, m, ha='center', va='center', fontsize=12, color='black')
            #ax.text(0.5, 0.6, f"", ha='right', va='center', fontsize=12, color='black')
            #ax.text(0.5, 0.6, f"", ha='right', va='center', fontsize=12, color='black')
            #ax.text(0.5, 0.6, f"", ha='right', va='center', fontsize=12, color='black')
 
        def clas(x):
            if 0.6<x<=1:return 0
            elif 0.4<x<=0.6:return 1
            elif 0.2<x<=0.4:return 2
            elif -1<=x<=0.2:return 3

        def cambio2(x,y):
            global factor
            factor=np.array([[None,-1 ,-4  ,-6],
                            [1   ,None,-2  ,-5],
                            [4   ,2   ,None,-3],
                            [6   ,5   ,3   ,None]])
            return factor[x][y]

        def listar_tif(x):
            return [i for i in x if i[-4:]==".tif"]
        
        def retransformar(x,rut):
            imgs=[]
            os.makedirs(os.path.join(rut,"img_origen_nacional"))
            for i in x:
                print(i)
                src=rs.open(os.path.join(ruta,i))
                transform, width, height = calculate_default_transform(src.crs, "EPSG:9377", src.width, src.height, *src.bounds)
                kwargs = src.meta.copy()
                kwargs.update({
                    'crs': "EPSG:9377",
                    'transform': transform,
                    'width': width,
                    'height': height
                })

            # Crear el nuevo archivo raster de salida
                dst=rs.open(rut+"/img_origen_nacional"+f"/{i}", 'w', **kwargs)
                for a in range(1, src.count + 1):
                    reproject(
                        source=rs.band(src, a),
                        destination=rs.band(dst, a),
                        src_transform=src.transform,
                        src_crs=src.crs,
                        dst_transform=transform,
                        dst_crs="EPSG:9377",
                        resampling=rs.enums.Resampling.nearest  # Cambia el método de interpolación según sea necesario
                    )
                dst.close()
                src.close()
                fn=rs.open(rut+"/img_origen_nacional"+f"/{i}")
                imgs.append(fn)
            return imgs

        def crear_carpeta(ruta):
            conteo=1
            if "resultado" in self.directorio:
                while True:
                    if f"resultado ({conteo})" not in self.directorio:
                        os.mkdir(os.path.join(ruta,f"resultado ({conteo})"))
                        return os.path.join(ruta,f"resultado ({conteo})")
                    conteo+=1
            else:
                os.mkdir(os.path.join(ruta,"resultado"))
                return os.path.join(ruta,"resultado")
        
        def guardar_ndvi(ruta,tiffs,data,crs,width,height,transform):
            print(tiffs)
            os.mkdir(os.path.join(ruta,"geotiff"))
            ruta=os.path.join(ruta,"geotiff")
            for i in range(len(data)):
                ndviImage = rs.open(os.path.join(ruta,f"NDVI_{tiffs[i]}"), 'w', driver='Gtiff',
                        width=width, height=height,
                        count=1,
                        crs=crs,
                        transform=transform,
                        dtype='float64'                  
                        )
                ndviImage.write(data[i],1) #ndvi
                ndviImage.close()
        
        def mostrar_ndvi_individual(ruta,names,ndvi,tipo):
            nombre=ruta.split("/")
            nombre=nombre[-1].split("\\")
            prom=[]
            años=[]
            #plt.ion()
            os.mkdir(os.path.join(ruta,"graficos"))
            for i in range(len(names)):  
                prom.append(round(np.nanmean(ndvi[i]),3))
                años.append(names[i][:-4])
                fig = plt.figure(figsize=(14,6))
                gs = fig.add_gridspec(1, 2, width_ratios=[2, 1]) 
                ax = fig.add_subplot(gs[0]) 
                ax2 = fig.add_subplot(gs[1]) 
                im = ax.imshow(ndvi[i], cmap='RdYlGn',vmax=1)
                plt.colorbar(im, ax=ax,label='NDVI')
                fig.suptitle(f'Índice de Vegetación de Diferencia Normalizada (NDVI) para {names[i][:-4]} en {nombre[0]}',weight='bold',size=16)
                ax.axis('off')
                info(ax2,ndvi[i],porcentaje(prom,años))
                ax2.set_xlim(0, 1)
                ax2.set_ylim(0, 1)
                ax2.set_xticks([])
                ax2.set_yticks([])

                inner_left = 0.6658
                inner_bottom = 0.1
                inner_width = 0.234#5
                inner_height = 0.3
                inner_ax = fig.add_axes([inner_left, inner_bottom, inner_width, inner_height])
                inner_ax.hist(ndvi[i].flatten(),label="NDVI",color="green",bins=200,range=(-1,1))   
                if tipo==1:
                    plt.get_current_fig_manager().window.state('zoomed')
                    plt.show()
                fig.savefig(os.path.join(os.path.join(ruta,f"graficos"),f"visual_NDVI_{names[i][:-4]}.png"), dpi=300,bbox_inches='tight')
                plt.close(fig)
            #plt.ion()
            global inicio
            inicio=datetime.now()
            for i in range(len(names)):
                fig, ax = plt.subplots()
                im = ax.imshow(ndvi[i], cmap='RdYlGn',vmax=1)
                fig.colorbar(im, ax=ax,label='NDVI')
                plt.axis('off')
                plt.ioff()
                plt.savefig(os.path.join(os.path.join(ruta,f"graficos"),f"grafico_NDVI_{names[i][:-4]}.png"), dpi=300,bbox_inches='tight')
                plt.close(fig)
                
                fig, ax = plt.subplots()
                ax.hist(ndvi[i].flatten(),label="NDVI",color="green",bins=200,range=(-1,1))
                plt.xlabel('NDVI')
                plt.ylabel('Frecuencia')
                plt.grid(True)
                plt.ioff()
                plt.savefig(os.path.join(os.path.join(ruta,f"graficos"),f"histograma_NDVI_{names[i][:-4]}.png"), dpi=300,bbox_inches='tight')
                plt.close(fig)
            
            print("Graficos guardados :D")
               
        def alertas(x,ruta,crs,transform):
            rt_detec=os.path.join(ruta,"deteccion")
            os.mkdir(rt_detec)
            matriz = np.empty((len(x[-1]),len(x[-1][0])))
            matriz[:] = np.nan
            c=0
            tip={0:"densa",1:"moderada",2:"escasa",3:"limpio"}
            ar=open(os.path.join(rt_detec,"registro.txt"),"w")
            for i in range(len(x[-1])):
                for e in range(len(x[-1][i])):
                    if not(np.isnan(x[-1][i][e])) and not(np.isnan(x[-2][i][e])):
                    #print(tip[clas(x[-1][i][e])],tip[clas(x[-2][i][e])])
                        if tip[clas(x[-1][i][e])]!=tip[clas(x[-2][i][e])]:
                            ar.write(f"{tip[clas(x[-1][i][e])]} - {tip[clas(x[-2][i][e])]} | {x[-1][i][e]} - {x[-2][i][e]}\n")
                            matriz[i][e]=cambio2(clas(x[-2][i][e]),clas(x[-1][i][e]))
                        c+=1
            ar.close()
            matriz=matriz.astype("float32")
            geometry = [shape(geom) for geom, value in rs.features.shapes(matriz, transform=transform) if not(np.isnan(value))]
            gdf = gpd.GeoDataFrame({'geometry': geometry},crs=crs)
            gdf['valor'] = [value for geom, value in rs.features.shapes(matriz, transform=transform) if not(np.isnan(value))]
            gdf['area'] = gdf['geometry'].area
            gdf['cambio']= [f"{tip[np.where(factor==int(value))[0][0]]} >> {tip[np.where(factor==int(value))[1][0]]}" for geom, value in rs.features.shapes(matriz, transform=transform) if not(np.isnan(value))]
            gdf['dscrpcn']=[f"Posible INCREMENTO de vegetacion" if int(value) in (4,7,8,10,11,12) else f"Posible DISMINUCION de vegetacion" for geom, value in rs.features.shapes(matriz, transform=transform) if not(np.isnan(value))]
            gdf = gdf[gdf['geometry'].area >= 37]
            gdf["x_cntro"]=gdf["geometry"].centroid.x
            gdf["y_cntro"]=gdf["geometry"].centroid.y
            lista=list(zip(gdf["x_cntro"],gdf["y_cntro"]))
            transformer = Transformer.from_crs('EPSG:9377','EPSG:4326',always_xy=True)
            lista = np.array(list(transformer.itransform(lista)))
            gdf['x']=lista[:,0]
            gdf['y']=lista[:,1]
            gdf.to_excel(os.path.join(rt_detec,"registro_deteccion.xlsx"),index=False, engine="openpyxl")
            
            gdf_menos=gdf[gdf['valor'].isin([1,2,3,5,6,9])]
            gdf_mas=gdf[gdf['valor'].isin([4,7,8,10,11,12])]


            carpeta_shp=os.path.join(rt_detec,"shp")
            os.mkdir(carpeta_shp)
            carpeta_kml=os.path.join(rt_detec,"kml")
            os.mkdir(carpeta_kml)

            gdf.to_file(os.path.join(carpeta_shp,"clasificacion.shp"))
            gpd.GeoDataFrame(geometry=gdf["geometry"].centroid,crs=crs).to_file(os.path.join(carpeta_shp,"puntos_clasificacion.shp"))
            
            gdf_menos.to_file(os.path.join(carpeta_shp,"posible_perdida.shp"))
            gpd.GeoDataFrame(geometry=gdf_menos["geometry"].centroid,crs=crs).to_file(os.path.join(carpeta_shp,"puntos_perdida.shp"))
            
            gdf_mas.to_file(os.path.join(carpeta_shp,"posible_incremento.shp"))
            gpd.GeoDataFrame(geometry=gdf_mas["geometry"].centroid,crs=crs).to_file(os.path.join(carpeta_shp,"puntos_incremento.shp"))
            print(f'Archivo shapefile creado en: {carpeta_shp}')
            
            shp=gpd.read_file(os.path.join(carpeta_shp,"clasificacion.shp"))                                                       
            gpd.io.file.fiona.drvsupport.supported_drivers['KML']='rw'  
            shp.to_file(os.path.join(carpeta_kml,"clasificacion.kml"),driver="KML") 
            shpp=gpd.read_file(os.path.join(carpeta_shp,"puntos_clasificacion.shp"))                                                       
            gpd.io.file.fiona.drvsupport.supported_drivers['KML']='rw'  
            shpp.to_file(os.path.join(carpeta_kml,"puntos_clasificacion.kml"),driver="KML") 
            
            shp2=gpd.read_file(os.path.join(carpeta_shp,"posible_perdida.shp"))                                                       
            gpd.io.file.fiona.drvsupport.supported_drivers['KML']='rw'  
            shp2.to_file(os.path.join(carpeta_kml,"posible_perdida.kml"),driver="KML")
            shpp2=gpd.read_file(os.path.join(carpeta_shp,"puntos_perdida.shp"))                                                       
            gpd.io.file.fiona.drvsupport.supported_drivers['KML']='rw'  
            shpp2.to_file(os.path.join(carpeta_kml,"puntos_perdida.kml"),driver="KML") 

            shp3=gpd.read_file(os.path.join(carpeta_shp,"posible_incremento.shp"))                                                       
            gpd.io.file.fiona.drvsupport.supported_drivers['KML']='rw'  
            shp3.to_file(os.path.join(carpeta_kml,"posible_incremento.kml"),driver="KML")
            shpp3=gpd.read_file(os.path.join(carpeta_shp,"puntos_incremento.shp"))                                                       
            gpd.io.file.fiona.drvsupport.supported_drivers['KML']='rw'  
            shpp3.to_file(os.path.join(carpeta_kml,"puntos_incremento.kml"),driver="KML") 

        def mostrar_img(x):
            pass
        #--------------------------------------------------
        
        
        if tipo==1:
            tifs=listar_tif(self.directorio)
            ruta_resultado=crear_carpeta(self.ruta)
            imgs=retransformar(tifs,ruta_resultado) #[rs.open(os.path.join(ruta,i)) for i in tifs]
            arrays=[i.read().astype('float64') for i in imgs]
            ndvis=[ndvi(i[2],i[3]) for i in arrays]
            guardar_ndvi(ruta_resultado,tifs,ndvis,imgs[0].crs,imgs[0].width,imgs[0].height,imgs[0].transform)
            mostrar_ndvi_individual(ruta_resultado,tifs,ndvis,1)
            alertas(ndvis,ruta_resultado,imgs[0].crs,imgs[0].transform)

        else:
            carpetas=self.directorio
            if "resultado" in self.directorio:
                carpetas=self.directorio[:self.directorio.index("resultado")]
                print(self.directorio)
            ruta_resultado=crear_carpeta(self.ruta)
            pro=1/len(self.directorio)
            for i in carpetas:
                print(f"\nTrabajando en {i}")
                ruta_c=os.path.join(ruta,i)
                tifs=listar_tif(os.listdir(ruta_c))
                imgs=[rs.open(os.path.join(ruta_c,i)) for i in tifs]
                arrays=[i.read().astype('float64') for i in imgs]
                ndvis=[ndvi(i[2],i[3]) for i in arrays]
                os.mkdir(os.path.join(ruta_resultado,i))
                guardar_ndvi(os.path.join(ruta_resultado,i),tifs,ndvis,imgs[0].crs,imgs[0].width,imgs[0].height,imgs[0].transform)
                mostrar_ndvi_individual(os.path.join(ruta_resultado,i),tifs,ndvis,0) 
                print("_______________________________________________") 
        ventana=ventana_secundaria()
        ventana.generado_correctamente_unico(ruta,f"tiempo total de ejecución: {(datetime.now()-inicio).total_seconds()} segundos")
        print(f"tiempo total: {(datetime.now()-inicio).total_seconds()} segundos | {(datetime.now()-inicio).total_seconds()*1000} milisegundos")
        
        """
        #--------------------------------------------------
        #lectura de imagenes

        img1=rs.open("./img/punto_1/planet2020.tif")
        img2=rs.open("./img/punto_1/img_sentinel2_2021.tif")

        #--------------------------------------------------
        #sentinel red=0 green=1 blue=2 nir=3
        #planet psscene blue=0 green=1 red=2 nir=3
        #extracción de bandas 
        array_img1=img1.read().astype('float64')
        img1_red=array_img1[2]
        img1_nir=array_img1[3]
        ndvi_img1=ndvi(img1_red,img1_nir)

        array_img2=img2.read().astype('float64')
        img2_red=array_img2[0]
        img2_nir=array_img2[3]
        ndvi_img2=ndvi(img2_red,img2_nir)


        #-------------------------------------------------- gist_earth gray viridis
        #mostrar imagenes bases
        fig,((ax1, ax2), (ax3, ax4))=plt.subplots(2,2)

        ax1.imshow(verdadero_color(array_img1[2],array_img1[1],array_img1[0]))
        ax1.axis("off")
        ax1.set_title("2020",fontdict={'weight':'bold'})

        ax2.imshow(verdadero_color(array_img2[0],array_img2[1],array_img2[2]))
        ax2.axis("off")
        ax2.set_title("2021",fontdict={'weight':'bold'})


        #mostrar ndvi
        plot.show(ndvi_img1,ax=ax3,cmap="RdYlGn",title="NDVI 2020")
        ax3.axis("off")
        plot.show(ndvi_img2,ax=ax4,cmap="RdYlGn",title="NDVI 2021")
        ax4.axis("off")

        plt.suptitle("Prueba piloto 01",weight='bold',size=16)
        plt.show()

        #---------------------------------------------------
        #histograma ndvi

        fig,((ax1, ax2), (ax3, ax4))=plt.subplots(2,2)
        plot.show(ndvi_img1,ax=ax1,cmap="RdYlGn",title="NDVI 2020")
        ax2.hist(ndvi_img1.flatten(),label="NDVI",color="green",bins=200)

        plot.show(ndvi_img2,ax=ax3,cmap="RdYlGn",title="NDVI 2021")
        ax4.hist(ndvi_img2.flatten(),label="NDVI",color="green",bins=200)

        plt.suptitle("Prueba piloto 01",weight='bold',size=16)
        plt.show()
        #---------------------------------------------------
        #ndvi individuales
        plt.figure(figsize=(10, 10))
        plt.imshow(ndvi_img1, cmap='RdYlGn',vmax=1)
        plt.colorbar(label='NDVI')
        plt.suptitle('Índice de Vegetación de Diferencia Normalizada (NDVI) para 2020',weight='bold',size=16)
        plt.axis('off')
        info(ndvi_img1)
        plt.subplots_adjust(right=0.83)
        plt.show()

        plt.figure(figsize=(10, 10))
        plt.imshow(ndvi_img2, cmap='RdYlGn',vmax=1)
        plt.colorbar(label='NDVI')
        plt.suptitle('Índice de Vegetación de Diferencia Normalizada (NDVI) para 2021',weight='bold',size=16)
        plt.axis('off')
        info(ndvi_img2)
        plt.subplots_adjust(right=0.83)
        plt.show()

        #---------------------------------------------------

        plt.figure(figsize=(10, 10))
        plt.imshow(clasificacion(ndvi_img2), cmap='RdYlGn',vmax=1)
        plt.colorbar(label='NDVI')
        plt.suptitle('Clasificación NDVI para 2020',weight='bold',size=16)
        plt.axis('off')
        plt.subplots_adjust(right=0.83)
        plt.show()


        #--------------------------------------------------

"""